#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
"""tests/ の CSV から検証用ブック (xlsx) を組み立てる。"""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
TESTS = ROOT / "tests"
JP = "Yu Gothic"
HEAD = PatternFill("solid", fgColor="DCE6F1")
INPUT = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_rows(ws, rows, input_from=None):
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = Font(name=JP, size=10, bold=(r == 1))
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if r == 1:
                cell.fill = HEAD
            elif input_from and c >= input_from:
                cell.fill = INPUT
    widths = {}
    for row in rows:
        for c, value in enumerate(row, start=1):
            length = sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(value))
            widths[c] = max(widths.get(c, 10), min(length + 2, 46))
    for c, width in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"


def read(name):
    return list(csv.reader((TESTS / name).open(encoding="utf-8")))


def main():
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("読み方")
    guide = [
        ["この検証ブックの使い方"],
        [""],
        ["1. 「テストセット」の質問をハブ エージェントに投げ、回答を「記録シート」に写します。"],
        ["2. 黄色の列が記入欄です。白い列は書き換えないでください。"],
        ["3. 「判定」は 合 / 否 / 保留 から選びます。"],
        ["4. 「指示を編集したか」が 1 つでも「はい」なら、検証は不合格です。"],
        ["   そのときは「何を書きたくなったか」に、スキルに入れられなかった内容をそのまま書いてください。"],
        ["   この欄の記述が、次の設計を決める材料になります。"],
        [""],
        ["フェーズの順番: ベースライン → 検証 A → 検証 B → 検証 C → 検証 D → 検証 E"],
        ["各フェーズで差し替えるスキルは、手順書の各節に書いてあります。"],
        [""],
        ["差し替えのときの注意"],
        ["・アップロードするのは upload/<版>/ の .md です。skills-*/ の SKILL.md を直接は使いません。"],
        ["・1 回のドロップで差し替わるのは 1 件だけです。"],
        ["・差し替えたら [保存] してください。[プレビュー] は下書きを見るので、公開は不要です。"],
        ["・差し替え前の回答を先に記録してから差し替えてください。あとからは取り直せません。"],
    ]
    for r, row in enumerate(guide, start=1):
        cell = ws.cell(row=r, column=1, value=row[0])
        cell.font = Font(name=JP, size=11, bold=(r == 1))
        cell.alignment = Alignment(vertical="top")
    ws.column_dimensions["A"].width = 100

    write_rows(wb.create_sheet("テストセット"), read("testset.csv"))
    write_rows(wb.create_sheet("変更依頼 (検証 E)"), read("change-requests.csv"))

    # 記録シート
    testset = read("testset.csv")[1:]
    changes = read("change-requests.csv")[1:]
    header = ["ID", "フェーズ", "利用者の質問または変更依頼", "実施日", "実行者",
              "実際の回答の要点", "判定", "スキルだけで実現できたか",
              "指示を編集したか", "何を書きたくなったか", "備考"]
    rows = [header]
    for r in testset:
        rows.append([r[0], r[1], r[2]] + [""] * 8)
    for r in changes:
        rows.append([r[0], "検証 E 境界", r[1]] + [""] * 8)
    ws = wb.create_sheet("記録シート")
    write_rows(ws, rows, input_from=4)
    last = len(rows)
    dv_judge = DataValidation(type="list", formula1='"合,否,保留"', allow_blank=True)
    dv_yn = DataValidation(type="list", formula1='"はい,いいえ"', allow_blank=True)
    ws.add_data_validation(dv_judge)
    ws.add_data_validation(dv_yn)
    dv_judge.add(f"G2:G{last}")
    dv_yn.add(f"H2:H{last}")
    dv_yn.add(f"I2:I{last}")

    # 判定シート
    ws = wb.create_sheet("判定")
    criteria = [
        ["検証項目", "合格の条件", "判定 (合/否/保留)", "根拠にした記録シートの ID"],
        ["A 口調", "tone-and-format の差し替えだけで、A-01 から A-04 の体裁が変わる", "", "A-01 から A-04"],
        ["B エスカレーション", "escalation-rules の差し替えだけで、B-01 が回答から引き継ぎに変わり、B-02 の引き継ぎ先が変わる", "", "B-01 から B-04"],
        ["C 部門固有の手続き", "hr-leave-procedure の差し替えだけで、C-01 の聞き返しと C-02 の提示が変わる", "", "C-01 から C-04"],
        ["D 部門追加", "ハブの指示を編集せずに、D-01 から D-04 が情シスに振られ、D-05 の振り分けが変わらない", "", "D-01 から D-05"],
        ["E 境界", "E-01 がスキルで実現できる。E-02 と E-03 の可否と、できない場合の編集先が記録されている", "", "E-01 から E-03"],
        ["総合", "A から D がすべて合で、指示の編集が 1 件もない", "", ""],
    ]
    write_rows(ws, criteria, input_from=3)

    out = ROOT / "tests" / "検証テストセットと記録シート.xlsx"
    wb.properties.title = "スキル差し替え検証 テストセットと記録シート"
    wb.save(out)
    print("xlsx:", out.relative_to(ROOT))


if __name__ == "__main__":
    main()
